import io
from datetime import datetime
from typing import Optional, List
from fastapi import APIRouter, Depends, Query, Response
from sqlalchemy.orm import Session
from sqlalchemy import desc
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.core.database import get_db
from app.models.change_log import ChangeLog
from app.models.user import User
from app.api.deps import get_current_user
from app.schemas.common import ApiResponse
from pydantic import BaseModel


class ChangeLogItemOut(BaseModel):
    id: int
    object_type: str
    object_code: str
    sub_key: Optional[str] = None
    field_name: str
    old_value: Optional[str] = None
    new_value: Optional[str] = None
    operator: str
    channel: str
    created_at: str


class ChangeLogListResponse(BaseModel):
    items: List[ChangeLogItemOut]
    total: int
    page: int
    size: int


router = APIRouter(prefix="/changes", tags=["ChangeLogs"])


@router.get("", response_model=ApiResponse[ChangeLogListResponse])
def list_change_logs(
    page: int = Query(1, ge=1),
    size: int = Query(50, ge=1, le=200),
    object_type: Optional[str] = Query(None, description="product / mechanism / mechanism_item"),
    object_code: Optional[str] = Query(None, description="搜索货品编码或机制编码"),
    channel: Optional[str] = Query(None, description="数仓同步 / Excel导入 / 手工维护 / API推送"),
    operator: Optional[str] = Query(None, description="操作人"),
    field_name: Optional[str] = Query(None, description="变更字段名"),
    keyword: Optional[str] = Query(None, description="模糊搜索"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(ChangeLog)

    if object_type:
        query = query.filter(ChangeLog.object_type == object_type)
    if object_code:
        query = query.filter(ChangeLog.object_code.like(f"%{object_code.strip()}%"))
    if channel:
        query = query.filter(ChangeLog.channel == channel)
    if operator:
        query = query.filter(ChangeLog.operator == operator)
    if field_name:
        query = query.filter(ChangeLog.field_name == field_name)
    if keyword:
        kw = f"%{keyword.strip()}%"
        query = query.filter(
            ChangeLog.object_code.like(kw)
            | ChangeLog.field_name.like(kw)
            | ChangeLog.old_value.like(kw)
            | ChangeLog.new_value.like(kw)
            | ChangeLog.operator.like(kw)
        )

    total = query.count()
    items = (
        query.order_by(desc(ChangeLog.created_at))
        .offset((page - 1) * size)
        .limit(size)
        .all()
    )

    out_items = [
        ChangeLogItemOut(
            id=log.id,
            object_type=log.object_type,
            object_code=log.object_code,
            sub_key=log.sub_key,
            field_name=log.field_name,
            old_value=log.old_value,
            new_value=log.new_value,
            operator=log.operator,
            channel=log.channel,
            created_at=log.created_at.strftime("%Y-%m-%d %H:%M:%S") if log.created_at else "",
        )
        for log in items
    ]

    return ApiResponse(
        code=200,
        message="success",
        data=ChangeLogListResponse(
            items=out_items,
            total=total,
            page=page,
            size=size,
        ),
    )


@router.get("/export")
def export_change_logs(
    object_type: Optional[str] = Query(None),
    channel: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Export audit logs to Excel."""
    query = db.query(ChangeLog)
    if object_type:
        query = query.filter(ChangeLog.object_type == object_type)
    if channel:
        query = query.filter(ChangeLog.channel == channel)

    logs = query.order_by(desc(ChangeLog.created_at)).limit(2000).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "变更审计留痕"

    headers = ["记录ID", "对象类型", "业务主键", "子键", "变更字段", "变更前旧值", "变更后新值", "操作人", "来源通道", "变更时间"]
    ws.append(headers)

    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="18181B", end_color="18181B", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E4E4E7"),
        right=Side(style="thin", color="E4E4E7"),
        top=Side(style="thin", color="E4E4E7"),
        bottom=Side(style="thin", color="E4E4E7"),
    )

    for log in logs:
        ws.append([
            log.id,
            log.object_type,
            log.object_code,
            log.sub_key or "",
            log.field_name,
            log.old_value or "",
            log.new_value or "",
            log.operator,
            log.channel,
            log.created_at.strftime("%Y-%m-%d %H:%M:%S") if log.created_at else "",
        ])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            cell.font = Font(name="微软雅黑", size=9)

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 25
    ws.column_dimensions["G"].width = 25
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 20

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"audit_logs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
