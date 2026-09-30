import io
from difflib import SequenceMatcher
from typing import List, Dict, Any, Optional, Tuple
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.product import Product
from app.models.mechanism import Mechanism
from app.schemas.quality import (
    QualityOverviewResponse,
    MissingFieldRecord,
    DuplicateCandidateRecord,
)


class QualityService:
    """Data Quality inspection, audit score calculation, and anomaly governance service."""

    @classmethod
    def get_overview(cls, db: Session) -> QualityOverviewResponse:
        total_p = db.query(Product).count()
        total_m = db.query(Mechanism).count()
        total_records = max(total_p + total_m, 1)

        missing_records = cls.get_missing_records(db, page=1, size=10000)[0]
        missing_count = len(missing_records)

        duplicate_candidates = cls.get_duplicate_candidates(db, threshold=0.85)
        duplicates_count = len(duplicate_candidates)

        pending_p = db.query(Product).filter(Product.pending_delete == True).count()
        pending_m = db.query(Mechanism).filter(Mechanism.pending_delete == True).count()
        pending_count = pending_p + pending_m

        # Dynamic health score calculation
        penalty = (missing_count * 1.5 + duplicates_count * 2.5 + pending_count * 2.0) / total_records * 10
        health_score = max(min(round(100.0 - penalty, 1), 100.0), 65.0)

        return QualityOverviewResponse(
            total_products=total_p,
            total_mechanisms=total_m,
            missing_fields_count=missing_count,
            duplicates_count=duplicates_count,
            pending_deletes_count=pending_count,
            health_score=health_score,
        )

    @classmethod
    def get_missing_records(
        cls,
        db: Session,
        domain: Optional[str] = None,
        page: int = 1,
        size: int = 50,
    ) -> Tuple[List[MissingFieldRecord], int]:
        all_records: List[MissingFieldRecord] = []

        if domain in (None, "product"):
            products = db.query(Product).all()
            for p in products:
                misses = []
                if not p.spec or str(p.spec).strip() == "":
                    misses.append("型号规格/净含量")
                if p.retail_price is None or p.retail_price <= 0:
                    misses.append("零售指导价")
                if not p.base_unit or str(p.base_unit).strip() == "":
                    misses.append("基本单位")
                if p.needs_maintenance == "是":
                    misses.append("运营待维护标记(需补充)")

                if misses:
                    all_records.append(
                        MissingFieldRecord(
                            domain="货品",
                            code=p.code,
                            name=p.name,
                            brand=p.brand,
                            missing_fields=misses,
                            data_source=p.data_source,
                            needs_maintenance=p.needs_maintenance or "否",
                        )
                    )

        if domain in (None, "mechanism"):
            mechanisms = db.query(Mechanism).all()
            for m in mechanisms:
                misses = []
                if not m.start_date:
                    misses.append("生效开始日期")
                if not m.end_date:
                    misses.append("生效结束日期")
                if m.mechanism_price is None or m.mechanism_price <= 0:
                    misses.append("机制核定结算价")
                if not m.items or len(m.items) == 0:
                    misses.append("无商品组合明细")

                if misses:
                    all_records.append(
                        MissingFieldRecord(
                            domain="机制",
                            code=m.code,
                            name=m.name,
                            brand=m.brand,
                            missing_fields=misses,
                            data_source=m.data_source,
                            needs_maintenance="是" if ("无商品组合明细" in misses) else "否",
                        )
                    )

        total = len(all_records)
        start_idx = (page - 1) * size
        end_idx = start_idx + size
        return all_records[start_idx:end_idx], total

    @classmethod
    def get_duplicate_candidates(
        cls,
        db: Session,
        brand: Optional[str] = None,
        threshold: float = 0.85,
    ) -> List[DuplicateCandidateRecord]:
        """Find suspected duplicate product records under the same brand using string similarity ratio."""
        query = db.query(Product).filter(Product.is_enabled == True)
        if brand:
            query = query.filter(Product.brand == brand)

        products = query.all()
        # Group by brand
        brand_map: Dict[str, List[Product]] = {}
        for p in products:
            b = p.brand or "未知品牌"
            brand_map.setdefault(b, []).append(p)

        candidates: List[DuplicateCandidateRecord] = []
        for b, plist in brand_map.items():
            if len(plist) < 2:
                continue
            for i in range(len(plist)):
                for j in range(i + 1, len(plist)):
                    p1 = plist[i]
                    p2 = plist[j]
                    if p1.code == p2.code:
                        continue

                    str1 = f"{p1.name} {p1.spec or ''}".strip().lower()
                    str2 = f"{p2.name} {p2.spec or ''}".strip().lower()
                    ratio = SequenceMatcher(None, str1, str2).ratio()

                    if ratio >= threshold:
                        candidates.append(
                            DuplicateCandidateRecord(
                                brand=b,
                                code_a=p1.code,
                                name_a=p1.name,
                                spec_a=p1.spec,
                                price_a=p1.retail_price,
                                status_a=p1.status or "正常",
                                code_b=p2.code,
                                name_b=p2.name,
                                spec_b=p2.spec,
                                price_b=p2.retail_price,
                                status_b=p2.status or "正常",
                                similarity_score=round(ratio, 2),
                            )
                        )

        candidates.sort(key=lambda x: x.similarity_score, reverse=True)
        return candidates

    @classmethod
    def export_missing_excel(cls, db: Session) -> io.BytesIO:
        records, _ = cls.get_missing_records(db, page=1, size=10000)
        wb = Workbook()
        ws = wb.active
        ws.title = "缺失关键属性清单"

        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0'),
        )

        headers = ["数据领域", "唯一编码", "名称", "所属品牌", "缺失关键属性项", "数据来源", "待维护标记", "排查建议"]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        data_font = Font(name="微软雅黑", size=10)
        for r in records:
            ws.append([
                r.domain,
                r.code,
                r.name,
                r.brand or "-",
                "、".join(r.missing_fields),
                r.data_source,
                r.needs_maintenance,
                "请运营同事通过 Excel 批量导入或界面详情补齐缺失字段",
            ])

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = data_font
                cell.border = border
                cell.alignment = Alignment(vertical="center")

        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 35
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 30
        ws.column_dimensions["F"].width = 16
        ws.column_dimensions["G"].width = 14
        ws.column_dimensions["H"].width = 36

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def export_pending_excel(cls, db: Session) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "待确认删除治理清单"

        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid")
        border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0'),
        )

        headers = ["数据领域", "唯一编码", "名称", "所属品牌", "数据来源", "状态", "失联原因", "最近更新时间"]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        data_font = Font(name="微软雅黑", size=10)
        products = db.query(Product).filter(Product.pending_delete == True).all()
        for p in products:
            ws.append([
                "货品",
                p.code,
                p.name,
                p.brand or "-",
                p.data_source,
                "启用" if p.is_enabled else "停用",
                "上游数仓同步连续缺失，触发待确认删除保护机制",
                p.source_updated_at.strftime("%Y-%m-%d %H:%M:%S") if p.source_updated_at else "-",
            ])

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = data_font
                cell.border = border
                cell.alignment = Alignment(vertical="center")

        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 35
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 12
        ws.column_dimensions["G"].width = 40
        ws.column_dimensions["H"].width = 22

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
